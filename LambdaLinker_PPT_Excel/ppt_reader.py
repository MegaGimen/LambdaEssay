from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Any, Optional

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except Exception as exc:  # pragma: no cover
    raise RuntimeError("未安装 python-pptx，请先运行：pip install python-pptx") from exc

try:
    from .text_utils import normalize_text
except ImportError:  # pragma: no cover
    from text_utils import normalize_text


@dataclass(frozen=True)
class PptDeckExtract:
    blocks: list[str]          # 给 LLM 的“内容块”
    meta: dict                 # deck 级别统计（触发视觉通道用）
    slide_meta: list[dict]     # slide 级别统计（未来做定向视觉补充用）


def _safe_int(x: Any, default: int = 0) -> int:
    try:
        return int(x)
    except Exception:
        return default


def read_pptx_deck(
    path: str,
    *,
    include_notes: bool = False,
    max_slides: Optional[int] = None,
    max_text_items_per_slide: int = 120,
) -> PptDeckExtract:
    """结构化读取 PPT：抽标题/文本框/备注；同时统计图片/图表/表格数量。"""
    if not path.lower().endswith(".pptx"):
        raise ValueError(f"仅支持 .pptx 文件: {path}")
    if not os.path.exists(path):
        raise FileNotFoundError(f"文件不存在: {path}")

    prs = Presentation(path)
    slides = list(prs.slides)
    if max_slides is not None:
        slides = slides[: max_slides]

    deck_blocks: list[str] = []
    slide_meta: list[dict] = []

    total_text_chars = 0
    total_text_boxes = 0
    total_pictures = 0
    total_charts = 0
    total_tables = 0

    for idx, slide in enumerate(slides, start=1):
        # 尽量按“视觉阅读顺序”排序（top/left）
        def _pos_key(shape: Any) -> tuple[int, int]:
            top = getattr(shape, "top", None)
            left = getattr(shape, "left", None)
            return (_safe_int(top, 10**12), _safe_int(left, 10**12))

        shapes_sorted = sorted(list(slide.shapes), key=_pos_key)

        title = ""
        try:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title = normalize_text(slide.shapes.title.text_frame.text or "")
        except Exception:
            title = ""

        parts: list[str] = [f"Slide {idx}"]
        if title:
            parts.append(f"Title: {title}")

        text_items: list[str] = []
        pictures = 0
        charts = 0
        tables = 0
        text_boxes = 0

        for shape in shapes_sorted:
            # --- 统计对象（用于触发视觉补充）---
            try:
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    pictures += 1
            except Exception:
                pass
            try:
                if getattr(shape, "has_chart", False):
                    charts += 1
            except Exception:
                pass
            try:
                if getattr(shape, "has_table", False):
                    tables += 1
            except Exception:
                pass

            # --- 抽文本 ---
            try:
                if not getattr(shape, "has_text_frame", False):
                    continue
                raw = shape.text_frame.text or ""
                txt = normalize_text(raw)
                if not txt:
                    continue
                if title and txt == title:
                    continue
                text_boxes += 1
                text_items.append(txt)
            except Exception:
                continue

        # 去重但保持顺序
        seen = set()
        uniq: list[str] = []
        for t in text_items:
            if t in seen:
                continue
            seen.add(t)
            uniq.append(t)

        uniq = uniq[: max_text_items_per_slide]
        if uniq:
            parts.append("Content:")
            for t in uniq:
                flat = " / ".join([line.strip() for line in t.splitlines() if line.strip()])
                parts.append(f"- {flat}")

        if include_notes:
            notes = ""
            try:
                notes = normalize_text(slide.notes_slide.notes_text_frame.text or "")
            except Exception:
                notes = ""
            if notes:
                flat = " / ".join([line.strip() for line in notes.splitlines() if line.strip()])
                parts.append("Notes:")
                parts.append(f"- {flat}")

        slide_text = "\n".join(uniq)
        slide_text_chars = len(slide_text)

        total_text_chars += slide_text_chars
        total_text_boxes += text_boxes
        total_pictures += pictures
        total_charts += charts
        total_tables += tables

        slide_meta.append(
            {
                "slide_index": idx,
                "title": title,
                "text_chars": slide_text_chars,
                "text_boxes": text_boxes,
                "pictures": pictures,
                "charts": charts,
                "tables": tables,
            }
        )
        deck_blocks.append("\n".join(parts).strip())

    meta = {
        "slides": len(deck_blocks),
        "total_text_chars": total_text_chars,
        "total_text_boxes": total_text_boxes,
        "total_pictures": total_pictures,
        "total_charts": total_charts,
        "total_tables": total_tables,
        "slides_with_nontext": sum(
            1 for m in slide_meta if (m.get("pictures", 0) or m.get("charts", 0) or m.get("tables", 0))
        ),
    }

    if not deck_blocks:
        deck_blocks = ["(Empty deck or no readable content)"]

    return PptDeckExtract(blocks=deck_blocks, meta=meta, slide_meta=slide_meta)
